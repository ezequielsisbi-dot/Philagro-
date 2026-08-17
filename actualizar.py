#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador del Dashboard de Cobranzas — Philagro S.A.

    python3 actualizar.py Cobranzas_2026.xlsx dashboard.html

Toma un Excel de cobranzas en formato contable y produce un HTML autocontenido
(se abre con doble clic, sin internet). Si el HTML de salida ya existe, se usa
como BASE: los datos nuevos se ANEXAN sobre el histórico sin perder ni duplicar.

Los criterios (moneda, signos, conceptos, contado/diferido) están congelados y
documentados en CLAUDE.md; este script sólo los aplica.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sys
import unicodedata
from collections import Counter
from typing import Any

RAIZ = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(RAIZ, "assets")

# Conceptos de forma de cobro, en el orden en que se les asigna color en el
# dashboard (criterio congelado #3).
FORMAS_ORDEN = [
    "Valores (echeqs/cheques)",
    "Bancos / transferencias",
    "Tarjetas / plataformas",
    "Caja Compensación",
    "Valores en recaudadoras",
    "Impuestos / retenciones",
    "Caja Chica",
    "Otros",
]

# Sólo estos dos conceptos pueden quedar DIFERIDOS (criterio congelado #9).
FORMAS_DIFERIBLES = ("Valores (echeqs/cheques)", "Valores en recaudadoras")

# Días entre el vencimiento del valor (Fechavto) y su acreditación efectiva.
DIAS_ACREDITACION = 2

MESES_ABBR = ["ene.", "feb.", "mar.", "abr.", "may.", "jun.",
              "jul.", "ago.", "sep.", "oct.", "nov.", "dic."]


# --------------------------------------------------------------------------- #
# Utilidades
# --------------------------------------------------------------------------- #
def normalizar(texto: Any) -> str:
    """Minúsculas, sin acentos y con espacios colapsados (para comparar textos)."""
    s = "" if texto is None else str(texto)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def a_fecha(valor: Any) -> dt.date | None:
    """Convierte a date lo que venga del Excel (datetime, date o texto)."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if isinstance(valor, (int, float)):  # serial de Excel (1899-12-30 base)
        try:
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(valor))).date()
        except (ValueError, OverflowError):
            return None
    txt = str(valor).strip()
    if not txt:
        return None
    for patron in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(txt[:10], patron).date()
        except ValueError:
            continue
    return None


def a_numero(valor: Any) -> float:
    """Convierte a float un importe del Excel. Los signos vienen del archivo:
    el Haber ya llega negativo y NO se lo toca (criterio congelado #2)."""
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    txt = str(valor).strip()
    if not txt:
        return 0.0
    negativo = txt.startswith("(") and txt.endswith(")")
    txt = txt.strip("()").replace("$", "").replace("US$", "").replace(" ", "")
    # Formato es-AR: el punto es separador de miles y la coma, decimal.
    if "," in txt:
        txt = txt.replace(".", "").replace(",", ".")
    elif txt.count(".") > 1:
        txt = txt.replace(".", "")
    try:
        n = float(txt)
    except ValueError:
        return 0.0
    return -n if negativo else n


def mes_de(fecha: dt.date) -> str:
    return f"{fecha.year:04d}-{fecha.month:02d}"


def etiqueta_mes(clave: str) -> str:
    anio, mes = clave.split("-")
    return f"{MESES_ABBR[int(mes) - 1]} {anio}"


# --------------------------------------------------------------------------- #
# Mapa de cuentas contables -> concepto
# --------------------------------------------------------------------------- #
class MapaCuentas:
    def __init__(self, ruta: str):
        with open(ruta, encoding="utf-8") as fh:
            cfg = json.load(fh)
        self.reglas = [(re.compile(r["patron"]), r["concepto"]) for r in cfg["reglas"]]
        self.por_defecto = cfg.get("por_defecto", "Otros")
        self.sin_mapear: Counter[str] = Counter()

    def concepto(self, cuenta: Any) -> str:
        nombre = normalizar(cuenta)
        for patron, concepto in self.reglas:
            if patron.search(nombre):
                return concepto
        self.sin_mapear[str(cuenta).strip()] += 1
        return self.por_defecto


# --------------------------------------------------------------------------- #
# Lectura del Excel
# --------------------------------------------------------------------------- #
# Alias de encabezado -> campo lógico. Se comparan normalizados.
ALIAS = {
    "fecha": ("fecha", "fecha cobranza", "fecha emision", "fecha comprobante",
              "fecha recibo", "fecha mov", "fecha movimiento"),
    "documento": ("documento", "comprobante", "nro documento", "nro comprobante",
                  "numero documento", "recibo", "nro recibo", "doc"),
    "cliente": ("cliente", "razon social", "nombre cliente", "nombre", "titular",
                "cliente razon social"),
    "cuenta": ("cuenta", "cuenta contable", "descripcion cuenta", "nombre cuenta",
               "cuenta descripcion"),
    "importe_sec": ("importe moneda secundaria", "imp moneda secundaria",
                    "importe moneda sec", "moneda secundaria", "importe usd",
                    "importe en dolares", "importe dolares", "importe u$s"),
    "importe": ("importe", "importe moneda principal", "imp moneda principal",
                "importe pesos", "importe ars", "monto", "importe original"),
    "fechavto": ("fechavto", "fecha vto", "fecha vencimiento", "vencimiento",
                 "fec vto", "fecha de vencimiento", "vto"),
    "debe": ("debe",),
    "haber": ("haber",),
}


def mapear_encabezado(celdas: list[Any]) -> dict[str, int]:
    """Devuelve {campo lógico: índice de columna} para una fila de encabezado."""
    cols: dict[str, int] = {}
    for i, celda in enumerate(celdas):
        nombre = normalizar(celda)
        if not nombre:
            continue
        for campo, alias in ALIAS.items():
            if campo in cols:
                continue
            if nombre in alias or any(nombre.startswith(a) for a in alias):
                cols[campo] = i
                break
    return cols


def leer_excel(ruta: str, mapa: MapaCuentas, verbose: bool = True) -> tuple[list[dict], dict]:
    """Lee el Excel y devuelve (movimientos, info). Un movimiento es una línea
    contable ya limpia; la cantidad de COBROS se cuenta después por Documento."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Falta openpyxl. Instalalo con:  pip install openpyxl")

    wb = load_workbook(ruta, read_only=True, data_only=True)
    hoja = wb[wb.sheetnames[0]]
    filas = list(hoja.iter_rows(values_only=True))
    wb.close()

    # Encabezado: primera fila que mapee al menos 3 campos conocidos.
    cols: dict[str, int] = {}
    fila_enc = -1
    for i, fila in enumerate(filas[:40]):
        candidato = mapear_encabezado(list(fila))
        if len(candidato) >= 3 and "fecha" in candidato:
            cols, fila_enc = candidato, i
            break
    if fila_enc < 0:
        muestra = [str(c) for c in (filas[0] if filas else [])][:20]
        sys.exit("No se encontró la fila de encabezados en el Excel.\n"
                 f"Primera fila leída: {muestra}")

    faltan = [c for c in ("fecha", "cuenta") if c not in cols]
    if faltan:
        sys.exit(f"El Excel no trae columna(s) obligatoria(s): {faltan}\n"
                 f"Columnas detectadas: {sorted(cols)}")

    # Moneda: USD si el archivo trae la columna real de moneda secundaria
    # (criterio congelado #1: nunca se inventa cotización).
    usa_usd = "importe_sec" in cols
    if not usa_usd and verbose:
        print("  ! El archivo NO trae 'Importe moneda secundaria': "
              "se usa el importe principal en ARS.")

    movs: list[dict] = []
    descartadas = 0
    for fila in filas[fila_enc + 1:]:
        val = lambda campo: (fila[cols[campo]] if campo in cols and cols[campo] < len(fila) else None)

        fecha = a_fecha(val("fecha"))
        if fecha is None:
            descartadas += 1
            continue

        if usa_usd:
            importe = a_numero(val("importe_sec"))
        elif "importe" in cols:
            importe = a_numero(val("importe"))
        else:  # sin columna de importe: se deriva de Debe/Haber
            importe = a_numero(val("debe")) - a_numero(val("haber"))

        cliente = str(val("cliente") or "SIN CLIENTE").strip() or "SIN CLIENTE"
        doc = str(val("documento") or "").strip()
        forma = mapa.concepto(val("cuenta"))
        fvto = a_fecha(val("fechavto"))

        movs.append({
            "fecha": fecha.isoformat(),
            "mes": mes_de(fecha),
            "cliente": cliente,
            "forma": forma,
            "importe": round(importe, 2),
            "doc": doc,
            # fa = fecha exacta de acreditación (Fechavto + 2 días). Se guarda para
            # poder recalcular contado/diferido cuando la fecha de corte avance.
            "fa": (fvto + dt.timedelta(days=DIAS_ACREDITACION)).isoformat() if fvto else None,
        })

    info = {"filas": len(filas), "fila_encabezado": fila_enc + 1, "descartadas": descartadas,
            "usa_usd": usa_usd, "columnas": sorted(cols)}
    return movs, info


# --------------------------------------------------------------------------- #
# Base: reconstrucción desde un dashboard ya generado
# --------------------------------------------------------------------------- #
def leer_base(ruta: str) -> tuple[list[dict], dict]:
    """Extrae REGISTROS y META de un dashboard HTML previo."""
    with open(ruta, encoding="utf-8") as fh:
        html = fh.read()

    def const(nombre: str) -> Any:
        m = re.search(r"^const\s+" + nombre + r"\s*=\s*(.+?);\s*$", html, re.M)
        if not m:
            sys.exit(f"El archivo base '{ruta}' no tiene 'const {nombre} = ...'. "
                     "¿Es un dashboard generado por este script?")
        return json.loads(m.group(1))

    return const("REGISTROS"), const("META")


# --------------------------------------------------------------------------- #
# Anexado (sin perder ni duplicar)
# --------------------------------------------------------------------------- #
def clave_sin_doc(r: dict) -> tuple:
    """Identidad por CONTENIDO para movimientos sin número de documento
    (ajustes tipo Caja Compensación)."""
    return (r["fecha"], r["cliente"], r["forma"], round(r["importe"], 2))


def anexar(base: list[dict], nuevos: list[dict]) -> tuple[list[dict], dict]:
    docs_base = {r["doc"] for r in base if r["doc"]}
    docs_nuevos = {r["doc"] for r in nuevos if r["doc"]}

    # Recibos CON número: si vienen en el archivo nuevo, mandan los del nuevo
    # (traen la Fechavto/acreditación exacta). Los demás del base se conservan.
    resultado = [r for r in base if not r["doc"] or r["doc"] not in docs_nuevos]

    # Movimientos SIN número: dedup por contenido, respetando repeticiones
    # legítimas (si el base ya tiene dos ajustes idénticos, se toleran dos).
    pendientes = Counter(clave_sin_doc(r) for r in base if not r["doc"])

    agregados_sin_doc = omitidos_sin_doc = 0
    for r in nuevos:
        if r["doc"]:
            resultado.append(r)
            continue
        k = clave_sin_doc(r)
        if pendientes[k] > 0:
            pendientes[k] -= 1
            omitidos_sin_doc += 1
            continue
        resultado.append(r)
        agregados_sin_doc += 1

    control = {
        "recibos_base": len(docs_base),
        "recibos_nuevo": len(docs_nuevos),
        "recibos_ya_estaban": len(docs_base & docs_nuevos),
        "recibos_agregados": len(docs_nuevos - docs_base),
        "movs_base": len(base),
        "movs_nuevo": len(nuevos),
        "ajustes_sin_doc_agregados": agregados_sin_doc,
        "ajustes_sin_doc_omitidos": omitidos_sin_doc,
    }
    return resultado, control


# --------------------------------------------------------------------------- #
# Contado / Diferido
# --------------------------------------------------------------------------- #
def marcar_contado_diferido(regs: list[dict], corte: str) -> None:
    """Fija 'a' (mes de acreditación) y 'd' (1 = diferido) en cada movimiento.

    DIFERIDO = valor (echeq / recaudadora) cuya acreditación es POSTERIOR a la
    fecha de corte = última cobranza del conjunto (criterio congelado #9). Se
    decide por fecha exacta cuando el movimiento trae Fechavto; los históricos
    sin Fechavto se aproximan por mes.
    """
    mes_corte = corte[:7]
    for r in regs:
        fa = r.get("fa")
        if r["forma"] not in FORMAS_DIFERIBLES:
            r["a"] = r["mes"]
            r["d"] = 0
        elif fa:
            r["a"] = fa[:7]
            r["d"] = 1 if fa > corte else 0
        else:
            # Histórico sin Fechavto: 'a' es lo que haya quedado guardado
            # (o el mes de cobro si nunca se supo).
            r["a"] = r.get("a") or r["mes"]
            if r["a"] > mes_corte:
                r["d"] = 1
            elif r["a"] < mes_corte:
                r["d"] = 0
            else:
                # Acredita dentro del mes de corte y no tenemos el día: se
                # respeta lo que ya se había decidido; si nunca se decidió,
                # se considera diferido.
                r["d"] = 1 if r.get("d", 1) else 0


# --------------------------------------------------------------------------- #
# Generación del HTML
# --------------------------------------------------------------------------- #
def generar_html(regs: list[dict], salida: str) -> dict:
    plantilla = os.path.join(ASSETS, "plantilla.html")
    chartjs = os.path.join(ASSETS, "chart.umd.min.js")
    for ruta in (plantilla, chartjs):
        if not os.path.exists(ruta):
            sys.exit(f"Falta {ruta}")

    regs = sorted(regs, key=lambda r: (r["fecha"], r["doc"], r["cliente"]))

    claves_mes = sorted({r["mes"] for r in regs} | {r["a"] for r in regs})
    meses = [{"key": k, "label": etiqueta_mes(k)} for k in claves_mes]
    formas = [f for f in FORMAS_ORDEN if any(r["forma"] == f for r in regs)]
    clientes = sorted({r["cliente"] for r in regs})
    fechas = [r["fecha"] for r in regs]
    docs = {r["doc"] for r in regs if r["doc"]}

    meta = {
        "moneda_codigo": MONEDA["codigo"],
        "moneda_nombre": MONEDA["nombre"],
        "generado": dt.date.today().strftime("%d/%m/%Y"),
        "total_cobros": len(docs) or len(regs),
        "total_movimientos": len(regs),
        "cobros_por_doc": bool(docs),
        "rango_min": min(fechas) if fechas else "",
        "rango_max": max(fechas) if fechas else "",
    }

    # En el HTML sólo viaja lo que el dashboard usa, más 'fa' (necesaria para
    # recalcular contado/diferido en la próxima actualización).
    livianos = [{"fecha": r["fecha"], "mes": r["mes"], "cliente": r["cliente"],
                 "forma": r["forma"], "importe": r["importe"], "doc": r["doc"],
                 "a": r["a"], "d": r["d"], **({"fa": r["fa"]} if r.get("fa") else {})}
                for r in regs]

    def dump(obj: Any) -> str:
        return json.dumps(obj, ensure_ascii=False).replace("</", "<\\/")

    html = open(plantilla, encoding="utf-8").read()
    reemplazos = {
        "__REGISTROS_JSON__": dump(livianos),
        "__MESES_JSON__": dump(meses),
        "__FORMAS_JSON__": dump(formas),
        "__CLIENTES_JSON__": dump(clientes),
        "__META_JSON__": dump(meta),
        "__CHARTJS__": open(chartjs, encoding="utf-8").read(),
    }
    for marca, valor in reemplazos.items():
        if marca not in html:
            sys.exit(f"La plantilla no tiene el placeholder {marca}")
        html = html.replace(marca, valor, 1)

    with open(salida, "w", encoding="utf-8") as fh:
        fh.write(html)
    return meta


MONEDA = {"codigo": "USD", "nombre": "dólares"}


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def plata(v: float) -> str:
    return f"{v:,.2f}".replace(",", "@").replace(".", ",").replace("@", ".")


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Genera el Dashboard de Cobranzas de Philagro a partir de un Excel.")
    ap.add_argument("excel", nargs="?",
                    help="Excel de cobranzas (Cobranzas_XXXX.xlsx). Se puede omitir "
                         "para re-generar el HTML desde el histórico, sin datos nuevos.")
    ap.add_argument("salida", nargs="?", default="dashboard.html",
                    help="HTML de salida (default: dashboard.html). Si ya existe, "
                         "se usa como base histórica y se sobrescribe.")
    ap.add_argument("--base", help="Dashboard HTML previo del que tomar el histórico "
                                   "(si no se indica, se usa el propio archivo de salida).")
    ap.add_argument("--sin-base", action="store_true",
                    help="Ignorar el histórico y generar sólo con este Excel.")
    ap.add_argument("--cuentas", action="store_true",
                    help="Sólo listar cómo quedó mapeada cada cuenta contable y salir.")
    args = ap.parse_args()

    # Si el primer posicional no es una planilla, es el HTML de salida:
    #   python3 actualizar.py --base dashboard.html nuevo.html
    if args.excel and not args.excel.lower().endswith((".xlsx", ".xlsm", ".xls")):
        if args.salida != "dashboard.html":
            sys.exit(f"'{args.excel}' no parece un Excel de cobranzas.")
        args.excel, args.salida = None, args.excel

    if args.excel and not os.path.exists(args.excel):
        sys.exit(f"No existe el archivo {args.excel}")

    mapa = MapaCuentas(os.path.join(ASSETS, "cuentas.json"))

    nuevos: list[dict] = []
    info = {"usa_usd": True}
    if args.excel:
        print(f"Leyendo {os.path.basename(args.excel)} …")
        nuevos, info = leer_excel(args.excel, mapa)
        print(f"  Encabezado en la fila {info['fila_encabezado']} · "
              f"{len(nuevos)} movimientos · columnas: {', '.join(info['columnas'])}")
        if info["descartadas"]:
            print(f"  {info['descartadas']} filas sin fecha descartadas (títulos, totales, vacías).")
    else:
        if args.sin_base or not (args.base or os.path.exists(args.salida)):
            sys.exit("Sin Excel hay que indicar un histórico: "
                     "python3 actualizar.py --base dashboard.html salida.html")
        print("Sin Excel nuevo: se re-genera el HTML desde el histórico.")

    if args.cuentas:
        print("\nCuentas sin mapear (irían a 'Otros'):")
        if not mapa.sin_mapear:
            print("  ninguna — todas las cuentas cayeron en un concepto conocido.")
        for cuenta, n in mapa.sin_mapear.most_common():
            print(f"  {n:>6}  {cuenta}")
        return

    if mapa.sin_mapear:
        print(f"\n  ATENCIÓN: {len(mapa.sin_mapear)} cuenta(s) sin mapear "
              f"({sum(mapa.sin_mapear.values())} movimientos) quedaron en 'Otros':")
        for cuenta, n in mapa.sin_mapear.most_common(15):
            print(f"    {n:>6}  {cuenta}")
        print("  Agregá el patrón en assets/cuentas.json y volvé a correr.\n")

    ruta_base = args.base or args.salida
    base: list[dict] = []
    meta_base: dict = {}
    if not args.sin_base and os.path.exists(ruta_base):
        base, meta_base = leer_base(ruta_base)
        print(f"Base: {os.path.basename(ruta_base)} · {len(base)} movimientos · "
              f"{meta_base.get('total_cobros', '?')} cobros")
    elif not args.sin_base:
        print("Base: no hay histórico previo, se arranca desde este Excel.")

    total_antes = sum(r["importe"] for r in base)
    regs, control = anexar(base, nuevos)

    corte = max((r["fecha"] for r in regs), default="")
    marcar_contado_diferido(regs, corte)
    meta = generar_html(regs, args.salida)
    total_despues = sum(r["importe"] for r in regs)

    cod = MONEDA["codigo"] if info["usa_usd"] else "ARS"
    print("\n--- Control de anexado ---------------------------------------")
    print(f"  Recibos en el base ............ {control['recibos_base']:>8}")
    print(f"  Recibos en el archivo nuevo ... {control['recibos_nuevo']:>8}")
    print(f"  Ya estaban (se reemplazaron) .. {control['recibos_ya_estaban']:>8}")
    print(f"  Se agregaron .................. {control['recibos_agregados']:>8}")
    if control["ajustes_sin_doc_agregados"] or control["ajustes_sin_doc_omitidos"]:
        print(f"  Ajustes sin documento ......... "
              f"{control['ajustes_sin_doc_agregados']} agregados, "
              f"{control['ajustes_sin_doc_omitidos']} ya estaban")
    print(f"  Movimientos ................... {control['movs_base']:>8} → {len(regs)}")
    print(f"  Total {cod} antes ............. {plata(total_antes):>16}")
    print(f"  Total {cod} después ........... {plata(total_despues):>16}")
    print(f"  Variación ..................... {plata(total_despues - total_antes):>16}")
    print(f"  Período ....................... {meta['rango_min']} a {meta['rango_max']}")
    print(f"  Fecha de corte (contado/dif.) . {corte}")
    dif = sum(r["importe"] for r in regs if r["d"])
    print(f"  Diferido al corte ............. {plata(dif):>16}")
    print("--------------------------------------------------------------")
    print(f"\nListo: {args.salida} ({os.path.getsize(args.salida) / 1e6:.1f} MB · "
          f"{meta['total_cobros']} cobros · {meta['total_movimientos']} movimientos)")


if __name__ == "__main__":
    main()
